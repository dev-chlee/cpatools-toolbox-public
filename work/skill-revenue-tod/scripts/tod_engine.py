"""
매출 TOD 엔진 - 프로파일 구동
==============================
샘플리스트 한 행(=장부에 기록된 매출=회계전표)을 외부 증빙으로 대사한다.
매출유형별 **검증 프로파일**(profiles.py)을 적용한다 — 해외/국내/용역 + 확장 가능.

구성 (결정형 = 그릇·룩업·롤업만; 판단은 스킬 구동 LLM 이 함):
  - 파일 열거·텍스트 아티팩트 로드·파일명 잠정분류(classify_file)
  - build_evidence_bundle(): LLM 판단용 thin 증빙 bundle 생성 (후보추출 없음 — raw text 제공)
  - analyze_sample_with_verdict(): LLM verdict 스키마검증 + 산술 재검증 + criticality 롤업
  - extract_samples_from_excel(): Claude/Codex가 확인한 mapping JSON으로 임의 시트 구조를 추출

사용법:
  python3 tod_engine.py <sample.xlsx> --evidence DIR --bundles-out bundles.json   # 판단재료 생성
  python3 tod_engine.py <sample.xlsx> --evidence DIR --verdicts verdicts.json --out results.json
"""

import sys
import openpyxl
import os
import json
import argparse
from openpyxl.utils import column_index_from_string

import cpa_storage
from incoterms import get_incoterms_policy


def _utf8_console():
    """Windows 기본 콘솔(cp949)에서 한글·기호(—, →) 출력이 깨지지 않도록 utf-8 재설정."""
    for _s in (sys.stdout, sys.stderr):
        if hasattr(_s, 'reconfigure'):
            try:
                _s.reconfigure(encoding='utf-8')
            except (AttributeError, OSError):
                pass


# ============================================================
# === 기본 경로 — 보관소(cpa_storage): CLI flag > CPA_SKILLS_STORAGE > ~/cpa-skills-data ===
# ============================================================

SKILL = "skill-revenue-tod"


def default_evidence_dir():
    """증빙 폴더 최상위 기본값: 보관소 input/evidence."""
    return str(cpa_storage.slot(SKILL, "input") / "evidence")


def default_output_json():
    """결과 JSON 기본값: 보관소 output/tod_detailed_results.json."""
    return str(cpa_storage.slot(SKILL, "output") / "tod_detailed_results.json")


# ========================
# 공통 유틸리티 함수 (결정형 재사용 코어)
# ========================

def extract_text_artifact(filepath):
    """
    사용자가 OCR/멀티모달 처리 후 제공한 텍스트 산출물을 읽는다.
    이 스킬은 PDF OCR이나 PDF 텍스트 추출을 직접 수행하지 않는다.
    """
    ext = filepath.lower().rsplit('.', 1)[-1] if '.' in filepath else ''
    with open(filepath, 'r', encoding='utf-8') as f:
        raw = f.read()
    if ext == 'json':
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            return ''
        return _text_from_json_artifact(data).strip()
    return raw.strip()


def _text_from_json_artifact(data):
    """Common OCR JSON shapes에서 본문만 추출한다. 메타데이터는 증빙 본문이 아니다."""
    if isinstance(data, str):
        return data
    if isinstance(data, dict):
        for key in ('full_text', 'text', 'markdown', 'md'):
            if isinstance(data.get(key), str):
                return data[key]
        content = data.get('content')
        if isinstance(content, dict):
            for key in ('full_text', 'text', 'markdown', 'html'):
                if isinstance(content.get(key), str):
                    return content[key]
        if isinstance(data.get('pages'), list):
            page_texts = [_text_from_json_artifact(p) for p in data['pages']]
            return "\n".join(t for t in page_texts if t)
        if isinstance(data.get('elements'), list):
            elem_texts = [_text_from_json_artifact(e) for e in data['elements']]
            return "\n".join(t for t in elem_texts if t)
        if isinstance(data.get('ocr_response'), dict):
            return _text_from_json_artifact(data['ocr_response'])
    if isinstance(data, list):
        item_texts = [_text_from_json_artifact(item) for item in data]
        return "\n".join(t for t in item_texts if t)
    return ''


# ========================
# 파일 분류 (파일명 기반 잠정 힌트 — 실제 문서 역할은 LLM 이 내용으로 확인)
# ========================

def classify_file(filename):
    """
    파일명 기반 증빙 분류.
    순서: 세금계산서 → 거래명세서 → 계약서 → POD → BL → CI → PO → INV/PL → OTHER
    (회계전표는 검증 '대상'(샘플)이므로 증빙으로 분류하지 않는다.)
    """
    fn = filename.lower()

    # 1. 세금계산서 (최우선)
    if '세금계산서' in fn:
        return '세금계산서'

    # 2. 거래명세서/인수증
    if any(x in fn for x in ['거래명세', '인수증']):
        return '거래명세서'

    # 3. 계약서 (용역매출 등) — 물류서류보다 앞
    if any(x in fn for x in ['계약서', '계약', 'contract', 'agreement']):
        return '계약서'

    # 4. POD
    if any(x in fn for x in ['pod', 'proof of delivery', 'delivery receipt',
                               '배송완료', '배송 서명', 'delivery order', 'pickup proof']):
        return 'POD'

    # 5. BL/AWB/출고증
    if any(x in fn for x in ['bl ', 'bl-', 'bl_', 'bl#', 'b/l', 'hawb', 'hbl', 'awb',
                               'air waybill', 'airwaybill', 'waybill', 'transportlabel',
                               'transport label', 'fedex', 'dhl', 'ups -', 'ups_',
                               '선적', 'skor', 'sur bl', 'sur_bl', 'sur ',
                               '출고증', '출고관련서류', '출고 증', 'kbe2', 'kbl0']):
        if 'pod' not in fn:
            return 'BL'

    # 6. CI (Commercial Invoice) - 단독 CI만
    if any(x in fn for x in ['ci-', 'ci_', 'commercial invoice']):
        if 'cipl' not in fn and 'ci&pl' not in fn:
            return 'CI'

    # 7. PO (Purchase Order)
    if any(x in fn for x in ['po ', 'po-', 'po_', 'po.', 'purchase order',
                               'p_us_', 'p_my_', 'p_uk_', 'p_de_', 'p_pt_', 'p_kor_',
                               'poimp', 'poinmt', 'pi -', 'pi_']):
        if 'pod' not in fn and 'invoice' not in fn and '출고' not in fn:
            return 'PO'

    # 8. INV/PL (Invoice/Packing List, 복합서류)
    if any(x in fn for x in ['inv_', 'inv-', 'inv ', 'inv0', 'invoice',
                               'cipl', 'ci&pl', 'pl&inv', '출고완료', '출고 완료']):
        return 'INV/PL'

    # 9. Order류
    if 'order' in fn and 'delivery' not in fn:
        return 'PO'

    return 'OTHER'


# ========================
# 증빙 폴더 로드·분류
# ========================

_ARTIFACT_EXTS = ('txt', 'md', 'markdown', 'json')


def _read_artifact_if_available(path):
    """읽기 실패도 미판독 증빙으로 남겨 다른 문서·샘플 처리를 계속한다."""
    try:
        return extract_text_artifact(path)
    except (OSError, UnicodeError):
        return ''


def _read_nested_artifact(folder):
    """문서 하위폴더 내부에서 텍스트 아티팩트(.md 우선)를 재귀 탐색해 로드."""
    for root, _dirs, files in os.walk(folder):
        arts = [f for f in files if f.lower().rsplit('.', 1)[-1] in _ARTIFACT_EXTS]
        if arts:
            arts.sort(key=lambda f: (0 if f.lower().endswith(('.md', '.markdown')) else 1, f))
            for artifact in arts:
                text = _read_artifact_if_available(os.path.join(root, artifact))
                if text.strip():
                    return text
    return ''


def _load_and_classify(folder_path):
    """
    샘플 폴더의 증빙을 분류·로드한다. {증빙타입: [doc_info]}.
      - flat 파일: 파일명으로 분류, 텍스트 아티팩트(.txt/.md/.json)면 로드.
      - 문서 하위폴더(<샘플>/<문서폴더>/<doc>.md): 폴더명으로 분류, 내부 아티팩트를 재귀 탐색해 로드.
    (분류는 파일명/폴더명 기반 잠정 힌트 — 실제 문서 역할은 LLM 이 내용으로 확인)
    """
    docs = {}
    for entry in sorted(os.listdir(folder_path)):
        epath = os.path.join(folder_path, entry)
        info = {'filename': entry, 'type': classify_file(entry), 'text': '', 'readable': False}

        if os.path.isdir(epath):
            text = _read_nested_artifact(epath)
            if text.strip():
                info['text'], info['readable'] = text, True
            else:
                info['text'] = '[문서폴더 내 OCR 텍스트(.txt/.md/.json) 없음]'
        elif os.path.isfile(epath):
            ext = entry.lower().rsplit('.', 1)[-1] if '.' in entry else ''
            if ext in _ARTIFACT_EXTS:
                text = _read_artifact_if_available(epath)
                if text.strip():
                    info['text'], info['readable'] = text, True
                else:
                    info['text'] = '[OCR 텍스트 본문 없음 또는 읽기 실패]'
            elif ext in ('pdf', 'jpg', 'jpeg', 'png'):
                info['text'] = '[원문 PDF/이미지 - OCR 텍스트 미제공]'
            else:
                info['text'] = '[지원하지 않는 형식 - OCR 텍스트(.txt/.md/.json) 필요]'
        else:
            continue

        docs.setdefault(info['type'], []).append(info)
    return docs


def _flatten_docs(docs):
    """Group dict -> thin LLM bundle list. 후보추출 없음 — LLM 이 raw text 를 읽어 판단한다.
    classified_type 은 파일명/폴더명 기반 **잠정** 분류(LLM 이 내용으로 확인)."""
    out = []
    idx = 1
    for doc_type in sorted(docs):
        for d in docs[doc_type]:
            out.append({
                'doc_id': f"D{idx:03d}",
                'filename': d.get('filename', ''),
                'classified_type': doc_type,          # 잠정 힌트
                'readable': bool(d.get('readable')),
                'text': d.get('text', ''),
            })
            idx += 1
    return out


def _docs_from_bundle(bundle):
    """LLM bundle docs list -> workpaper-compatible grouped dict."""
    grouped = {}
    for d in bundle.get('docs', []):
        doc_type = d.get('classified_type', 'OTHER')
        grouped.setdefault(doc_type, []).append({
            'filename': d.get('filename', ''),
            'type': doc_type,
            'text': d.get('text', ''),
            'readable': bool(d.get('readable')),
        })
    return grouped


# ========================
# 폴더 경로 결정
# ========================

def get_folder_path(sample, base=None):
    """기본 계약인 evidence_root/<folder> 경로를 반환한다."""
    base = base if base is not None else default_evidence_dir()
    folder = sample.get('folder', '')
    return os.path.join(base, folder)


def resolve_folder(sample, base):
    """일반 경로를 먼저 보고, 하위 트리의 고유한 동일 폴더명까지 지원한다."""
    folder = sample.get('folder', '')
    candidates = [
        get_folder_path(sample, base),
        os.path.join(base, sample.get('type', ''), folder),
    ]
    for candidate in candidates:
        if candidate and os.path.isdir(candidate):
            return candidate
    matches = []
    if os.path.isdir(base):
        for root, dirs, _files in os.walk(base):
            for dirname in dirs:
                if dirname == folder:
                    matches.append(os.path.join(root, dirname))
            if len(matches) > 1:
                break
    return matches[0] if len(matches) == 1 else candidates[0]


def _folder_ready(folder_path):
    """증빙 폴더가 얼라인되어 있는지: 존재 + 항목(파일 또는 문서 하위폴더) 1개 이상."""
    return bool(folder_path and os.path.isdir(folder_path) and os.listdir(folder_path))


def preflight_alignment(samples, evidence_root):
    """
    증빙 폴더가 준비되지 않은(얼라인 안 된) 샘플 목록을 반환한다.
    검증 전에 사용자에게 '증빙을 샘플당 폴더로 정리하라'고 알리기 위한 점검.
    반환: [{folder, type, expected_path, reason}]  (reason: '폴더 없음' | '폴더 비어있음')
    """
    unaligned = []
    for s in samples:
        folder = resolve_folder(s, evidence_root)
        if not _folder_ready(folder):
            reason = '폴더 비어있음' if (folder and os.path.isdir(folder)) else '폴더 없음'
            unaligned.append({'folder': s.get('folder', ''), 'type': s.get('type', ''),
                              'expected_path': folder, 'reason': reason})
    return unaligned


# ========================
# LLM-in-the-loop bundle / verdict 경로
# ========================

def _select_profile_id(sample, profile=None):
    """Return (profile_id, did_fallback)."""
    from profiles import PROFILES, select_profile

    pid = profile or select_profile(sample)
    did_fallback = pid not in PROFILES
    if did_fallback:
        pid = 'overseas'
    return pid, did_fallback


def build_evidence_bundle(sample, folder_path, profile=None):
    """
    Build deterministic input for LLM judgment.
    This function extracts and structures evidence; it does not judge.
    """
    from profiles import get_profile_spec

    pid, did_fallback = _select_profile_id(sample, profile)
    warnings = []
    if did_fallback:
        warnings.append(
            f"프로파일 자동판정 실패(매출유형='{sample.get('type', '')}') - {pid} 기본 적용")

    docs = {}
    if _folder_ready(folder_path):
        docs = _load_and_classify(folder_path)
    else:
        warnings.append('증빙 폴더 미정렬 — 샘플당 폴더로 정리 필요(폴더명=샘플 folder)')

    return {
        'sample': {
            'folder': sample.get('folder', ''),
            'type': sample.get('type', ''),
            'customer': sample.get('customer', ''),
            'krw_amount': sample.get('krw_amount'),
            'fc_amount': sample.get('fc_amount'),
            'currency': sample.get('currency', ''),
            'incoterms': sample.get('incoterms', ''),
            'incoterms_policy': get_incoterms_policy(sample.get('incoterms', '')),
            'ci_no': sample.get('ci_no', ''),
            'account_date': sample.get('account_date', ''),
            'voucher_no': sample.get('voucher_no', ''),
            'period': sample.get('period', ''),
        },
        'profile': get_profile_spec(pid),
        'docs': _flatten_docs(docs),
        'warnings': warnings,
    }


def _base_result(sample, profile_spec, docs=None):
    return {
        'folder': sample.get('folder', ''),
        'type': sample.get('type', ''),
        'profile': profile_spec['id'],
        'profile_display': profile_spec['display'],
        'base': {
            'ci_no_excel': sample.get('ci_no', ''),
            'customer_excel': sample.get('customer', ''),
            'incoterms_excel': sample.get('incoterms', ''),
            'currency_excel': sample.get('currency', ''),
            'fc_amount_excel': sample.get('fc_amount'),
            'krw_amount_excel': sample.get('krw_amount'),
            'account_date_excel': sample.get('account_date', ''),
            'voucher_no': sample.get('voucher_no', ''),
            'period': sample.get('period', ''),
        },
        'docs': docs or {},
        'checks': {},
        'overall': 'Pass',
        'exceptions': [],
    }


def _invalid_verdict_result(sample, bundle, errors):
    result = _base_result(sample, bundle['profile'], docs=_docs_from_bundle(bundle))
    result['overall'] = 'Exception'
    result['exceptions'] = [f"LLM verdict schema 오류: {e}" for e in errors]
    for c in bundle['profile']['checks']:
        result['checks'][c['id']] = {
            'label': c['label'],
            'result': 'Exception',
            'detail': 'LLM verdict schema 오류 - 수동검토 필요',
        }
    return result


def _missing_verdict_result(sample, bundle):
    result = _base_result(sample, bundle['profile'], docs=_docs_from_bundle(bundle))
    result['overall'] = 'Exception'
    result['exceptions'] = ['LLM verdict 누락 - 판단 결과 JSON 필요']
    for c in bundle['profile']['checks']:
        result['checks'][c['id']] = {
            'label': c['label'],
            'result': 'Exception',
            'detail': 'LLM verdict 누락 - 수동검토 필요',
        }
    return result


def _apply_evidence_gate(result, bundle):
    """증빙 부재·전부 미판독은 제출된 verdict보다 우선한다."""
    docs = bundle.get('docs', [])
    if not docs:
        status = 'Fail'
        reason = '증빙 없음 — 증빙 폴더가 없거나 비어 있어 검증할 수 없습니다.'
    elif not any(d.get('readable') and d.get('text', '').strip() for d in docs):
        status = 'Exception'
        reason = '증빙 미판독 — 읽을 수 있는 OCR 텍스트가 없어 수동확인이 필요합니다.'
    else:
        return result

    if result['overall'] != 'Fail':
        result['overall'] = status
    result['exceptions'].append(reason)
    for check in result['checks'].values():
        if check['result'] == 'Fail':
            check['detail'] += '\n' + reason
        else:
            check['result'] = status
            # 근거를 읽지 못한 Pass 설명을 최종 조서에 그대로 남기지 않는다.
            check['detail'] = reason
    return result


def analyze_sample_with_verdict(sample, folder_path, verdict, profile=None):
    """현재 Claude/Codex 대화에서 작성한 판정을 검증하여 조서 결과를 반환한다.

    LLM verdict를 스키마검증 → **금액 산술 재검증**(Python authority) → criticality 롤업 순으로 처리한다.
    """
    from llm_schema import (
        format_check_detail, reconcile_amounts, rollup_verdict,
        validate_evidence_citations, validate_verdict,
    )

    bundle = build_evidence_bundle(sample, folder_path, profile=profile)
    if verdict is None:
        return _apply_evidence_gate(_missing_verdict_result(sample, bundle), bundle)

    errors = validate_verdict(verdict, bundle['profile'])
    if any(doc.get('readable') for doc in bundle.get('docs', [])):
        errors += validate_evidence_citations(verdict, bundle.get('docs', []))
    if errors:
        return _apply_evidence_gate(_invalid_verdict_result(sample, bundle, errors), bundle)

    # 금액 authority = Python: LLM 지목값을 샘플과 재계산 대조(불일치 시 verdict를 Exception으로 강등)
    reconcile_amounts(verdict, sample, bundle['profile'])
    rolled = rollup_verdict(verdict, bundle['profile'])
    result = _base_result(sample, bundle['profile'], docs=_docs_from_bundle(bundle))
    result['overall'] = rolled['overall']
    result['exceptions'] = list(bundle.get('warnings', [])) + rolled['exceptions']

    for c in bundle['profile']['checks']:
        check = verdict['checks'][c['id']]
        result['checks'][c['id']] = {
            'label': check.get('label') or c['label'],
            'result': check['result'],
            'detail': format_check_detail(check),
        }
    return _apply_evidence_gate(result, bundle)


def load_verdicts(path):
    """Load verdict JSON as {sample_folder: verdict}."""
    with open(path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    if isinstance(data, dict) and isinstance(data.get('verdicts'), list):
        items = data['verdicts']
    elif isinstance(data, list):
        items = data
    elif isinstance(data, dict) and 'checks' in data:
        items = [data]
    elif isinstance(data, dict):
        out = {}
        for raw_key, verdict in data.items():
            key = str(raw_key)
            if not isinstance(verdict, dict):
                raise ValueError(f'verdict {key!r}는 JSON 객체여야 합니다.')
            declared = verdict.get('sample_folder') or verdict.get('folder')
            if declared is not None and str(declared) != key:
                raise ValueError(
                    f'verdict key {key!r}와 sample_folder {declared!r}가 일치하지 않습니다.'
                )
            verdict = dict(verdict)
            verdict.setdefault('sample_folder', key)
            out[key] = verdict
        return out
    else:
        raise ValueError('verdicts JSON은 배열, 단일 verdict 또는 {"verdicts":[...]} 형식이어야 합니다.')

    out = {}
    for v in items:
        if not isinstance(v, dict):
            raise ValueError('각 verdict는 JSON 객체여야 합니다.')
        key = v.get('sample_folder') or v.get('folder')
        if not key:
            raise ValueError('각 verdict에는 sample_folder가 필요합니다.')
        key = str(key)
        if key in out:
            raise ValueError(f'verdict sample_folder가 중복되었습니다: {key}')
        out[key] = v
    return out


# ============================================================
# === 샘플 추출 (Claude/Codex가 작성한 범용 mapping 사용) ===
# ============================================================

_REQUIRED_SAMPLE_FIELDS = {'folder', 'type', 'customer', 'krw_amount'}
_TEXT_SAMPLE_FIELDS = {
    'folder', 'type', 'customer', 'ci_no', 'incoterms', 'currency',
    'voucher_no', 'period',
}
_NUMERIC_SAMPLE_FIELDS = {'krw_amount', 'fc_amount'}


class SampleMappingError(ValueError):
    """샘플 목록은 있으나 안전하게 해석할 mapping이 없거나 잘못된 경우."""


def _workbook_has_values(wb):
    return any(
        cell.value not in (None, '')
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
    )


def _column_number(value):
    if isinstance(value, int) and value >= 1:
        return value
    if isinstance(value, str) and value.strip():
        try:
            return column_index_from_string(value.strip().upper())
        except ValueError as exc:
            raise SampleMappingError(f'잘못된 열 지정: {value!r}') from exc
    raise SampleMappingError(f'열은 A 같은 문자 또는 1 이상의 정수여야 합니다: {value!r}')


def load_sample_mapping(path):
    """UTF-8 JSON mapping을 읽는다. 형식은 references/sample-mapping.md 참조."""
    try:
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise SampleMappingError(f'mapping JSON을 읽을 수 없습니다: {exc}') from exc
    if isinstance(data, list):
        data = {'sections': data}
    if not isinstance(data, dict) or not isinstance(data.get('sections'), list):
        raise SampleMappingError('mapping 최상위에 sections 배열이 필요합니다.')
    if not data['sections']:
        raise SampleMappingError('mapping sections가 비어 있습니다.')
    return data


def extract_samples_from_excel(excel_path, mapping=None):
    """
    엑셀에서 샘플 데이터를 추출한다.

    mapping은 sections 배열이며 각 section은 sheet/start_row/end_row/columns/constants를
    가진다. Claude/Codex가 원본 통합문서를 읽어 열 의미를 확인한 뒤 작성한다.
    데이터가 있는 통합문서에 mapping이 없으면 0건 성공 대신 SampleMappingError를 낸다.

    반환: list of dict, 각 dict 권장키:
      필수:  folder(폴더명), type(매출유형), customer(거래처명), krw_amount(원화금액)
      해외:  ci_no, ci_date, incoterms, currency, fc_amount
      국내:  account_date, voucher_no
    type 값은 profiles.TYPE_TO_PROFILE 의 키와 맞추면 프로파일이 자동 선택된다
    (국내매출 / 중계무역 / 직수출 / 수출 / 해외매출 / 용역매출 / 용역).
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    try:
        has_values = _workbook_has_values(wb)
        if mapping is None:
            if has_values:
                raise SampleMappingError(
                    '표본 데이터가 있는 통합문서에는 --mapping이 필요합니다. '
                    'Claude/Codex가 시트·행·열을 확인해 mapping JSON을 작성하세요.'
                )
            return []

        sections = mapping.get('sections') if isinstance(mapping, dict) else None
        if not isinstance(sections, list) or not sections:
            raise SampleMappingError('mapping sections가 비어 있거나 올바르지 않습니다.')

        samples = []
        for section_no, section in enumerate(sections, 1):
            if not isinstance(section, dict):
                raise SampleMappingError(f'sections[{section_no}]는 객체여야 합니다.')
            sheet_name = section.get('sheet')
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    raise SampleMappingError(f'시트를 찾을 수 없습니다: {sheet_name!r}')
                ws = wb[sheet_name]
            else:
                ws = wb.active
            try:
                start_row = int(section['start_row'])
                end_row = int(section.get('end_row') or ws.max_row)
            except (KeyError, TypeError, ValueError) as exc:
                raise SampleMappingError(
                    f'sections[{section_no}]에 올바른 start_row/end_row가 필요합니다.'
                ) from exc
            if start_row < 1 or end_row < start_row:
                raise SampleMappingError(f'sections[{section_no}]의 행 범위가 잘못되었습니다.')

            columns = section.get('columns', {})
            constants = section.get('constants', {})
            if not isinstance(columns, dict) or not isinstance(constants, dict):
                raise SampleMappingError('columns와 constants는 객체여야 합니다.')
            unknown_overlap = set(columns) & set(constants)
            if unknown_overlap:
                raise SampleMappingError(
                    f'같은 필드를 columns와 constants에 중복 지정했습니다: {sorted(unknown_overlap)}'
                )
            missing = _REQUIRED_SAMPLE_FIELDS - (set(columns) | set(constants))
            if missing:
                raise SampleMappingError(f'필수 mapping 필드가 없습니다: {sorted(missing)}')
            column_numbers = {field: _column_number(col) for field, col in columns.items()}

            for row_no in range(start_row, end_row + 1):
                values = dict(constants)
                values.update({field: ws.cell(row_no, col).value
                               for field, col in column_numbers.items()})
                if values.get('folder') in (None, ''):
                    continue
                for field in _TEXT_SAMPLE_FIELDS:
                    if field in values and values[field] is not None:
                        values[field] = str(values[field]).strip()
                if any(values.get(field) in (None, '') for field in _REQUIRED_SAMPLE_FIELDS):
                    missing_row = [field for field in sorted(_REQUIRED_SAMPLE_FIELDS)
                                   if values.get(field) in (None, '')]
                    raise SampleMappingError(
                        f'{ws.title}!{row_no}행 필수값이 비어 있습니다: {missing_row}'
                    )
                for field in _NUMERIC_SAMPLE_FIELDS:
                    value = values.get(field)
                    if value in (None, '') and field != 'krw_amount':
                        continue
                    if isinstance(value, bool) or not isinstance(value, (int, float)):
                        raise SampleMappingError(
                            f'{ws.title}!{row_no}행 {field}가 숫자가 아닙니다: {value!r}'
                        )
                values['_source'] = {'sheet': ws.title, 'row': row_no}
                samples.append(values)

        if not samples and has_values:
            raise SampleMappingError('mapping 범위에서 표본을 한 건도 추출하지 못했습니다.')
        seen = set()
        duplicates = set()
        for sample in samples:
            folder = sample['folder']
            if folder in seen:
                duplicates.add(folder)
            seen.add(folder)
        duplicates = sorted(duplicates)
        if duplicates:
            raise SampleMappingError(f'표본 folder 값이 중복되었습니다: {duplicates}')
        return samples
    finally:
        wb.close()


# ============================================================
# MAIN
# ============================================================

def main(argv=None):
    _utf8_console()
    ap = argparse.ArgumentParser(
        description="매출 TOD 엔진 — 샘플리스트(전표)를 증빙으로 대사하여 결과 JSON 생성.")
    ap.add_argument('sample_xlsx', help='샘플리스트 엑셀 경로')
    ap.add_argument('--evidence',
                    help='증빙 폴더 최상위 (기본: 보관소 input/evidence)')
    ap.add_argument('--out',
                    help='결과 JSON 경로 (기본: 보관소 output/tod_detailed_results.json)')
    ap.add_argument('--profile', default='auto',
                    choices=['auto', 'overseas', 'domestic', 'service'],
                    help='프로파일 지정 (기본 auto = 샘플 매출유형으로 자동 선택)')
    ap.add_argument('--mapping',
                    help='샘플 엑셀의 시트·행·열 mapping JSON (비어 있지 않은 입력은 필수)')
    ap.add_argument('--bundles-out',
                    help='Claude/Codex 판단용 evidence bundle JSON 저장 경로')
    ap.add_argument('--verdicts',
                    help='Claude/Codex verdict JSON 경로. 지정 시 판정을 검증·롤업한다.')
    args = ap.parse_args(argv)
    try:
        args.evidence = args.evidence or default_evidence_dir()
        if args.verdicts:
            args.out = str(cpa_storage.guard(args.out)) if args.out else default_output_json()
        if args.bundles_out:
            cpa_storage.guard(args.bundles_out)
    except cpa_storage.StorageError as exc:
        ap.error(str(exc))

    try:
        mapping = load_sample_mapping(args.mapping) if args.mapping else None
        samples = extract_samples_from_excel(args.sample_xlsx, mapping=mapping)
    except SampleMappingError as exc:
        ap.error(str(exc))
    print(f"샘플 {len(samples)}건 추출")

    # 얼라인 점검: 증빙 폴더가 준비 안 된 샘플을 먼저 알린다 (얼라인은 사용자 책임)
    unaligned = preflight_alignment(samples, args.evidence)
    if unaligned:
        print(f"\n[얼라인 필요] 증빙 폴더가 준비되지 않은 샘플 {len(unaligned)}건:")
        for u in unaligned:
            print(f"  - {u['folder']} ({u['type']}): {u['reason']} → {u['expected_path']}")
        print("  증빙을 샘플당 폴더로 정리(폴더명 = 샘플 folder)한 뒤 다시 실행하세요.")
        print("  (처리는 계속하되, 위 샘플은 Fail 로 표기됩니다.)\n")

    if not args.bundles_out and not args.verdicts:
        ap.error("--bundles-out (판단재료 생성) 또는 --verdicts (Claude/Codex 판정 검증·롤업) 중 하나가 필요합니다.")

    forced = None if args.profile == 'auto' else args.profile
    try:
        verdict_map = load_verdicts(args.verdicts) if args.verdicts else None
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        ap.error(f'verdict JSON을 읽을 수 없습니다: {exc}')
    if verdict_map is not None:
        expected_folders = {sample['folder'] for sample in samples}
        unexpected = sorted(set(verdict_map) - expected_folders)
        if unexpected:
            ap.error(f'샘플 목록에 없는 verdict가 있습니다: {unexpected}')
    bundles = []
    all_results = []
    for i, sample in enumerate(samples):
        folder = resolve_folder(sample, args.evidence)
        print(f"[{i + 1}/{len(samples)}] {sample.get('folder', '')}...", end=' ')
        bundle = build_evidence_bundle(sample, folder, profile=forced)
        bundles.append(bundle)
        if verdict_map is not None:
            verdict = verdict_map.get(sample.get('folder', ''))
            r = analyze_sample_with_verdict(sample, folder, verdict, profile=forced)
            all_results.append(r)
            print(f"{r['profile_display']} → {r['overall']} (예외 {len(r['exceptions'])}건)")
        else:
            print(f"{bundle['profile']['display']} bundle ({len(bundle['docs'])} docs)")

    if args.bundles_out:
        os.makedirs(os.path.dirname(args.bundles_out) or '.', exist_ok=True)
        with open(args.bundles_out, 'w', encoding='utf-8') as f:
            json.dump({'bundles': bundles}, f, ensure_ascii=False, indent=2, default=str)
        print(f"\nClaude/Codex 검토 bundle 저장: {args.bundles_out}")

    if verdict_map is not None:
        os.makedirs(os.path.dirname(args.out) or '.', exist_ok=True)
        with open(args.out, 'w', encoding='utf-8') as f:
            json.dump(all_results, f, ensure_ascii=False, indent=2, default=str)
        p = sum(1 for r in all_results if r['overall'] == 'Pass')
        e = sum(1 for r in all_results if r['overall'] == 'Exception')
        fail = sum(1 for r in all_results if r['overall'] == 'Fail')
        print(f"\n=== SUMMARY: Pass={p}, Exception={e}, Fail={fail} (Total={len(all_results)}) ===")
        print(f"결과 저장: {args.out}")


if __name__ == '__main__':
    main()
